#!/usr/bin/env python3
"""Build Eli Margolin's master resume from one content block.

Outputs (all from the same CONTENT):
  ~/Desktop/Margolin_Resume.docx                  (python-docx, Calibri)
  ~/Desktop/Margolin_Resume.pdf                   (fpdf2, Arial)
  <site>/assets/files/Margolin_Resume.pdf         (copy of the PDF)
  <site>/resume/index.md                          (Jekyll `list` layout front matter)

Usage:  python3 _tools/build_resume.py            (run from the site repo root or anywhere)
Wording comes from the memory file canonical-resume-bullets.md; edit CONTENT here
and re-run rather than hand-editing the outputs.
"""
import os
import re
import shutil
import sys

HOME = os.path.expanduser("~")
DESKTOP = os.path.join(HOME, "Desktop")
SITE = os.path.join(DESKTOP, "ecmargo.github.io")
DOCX_OUT = os.path.join(DESKTOP, "Margolin_Resume.docx")
PDF_OUT = os.path.join(DESKTOP, "Margolin_Resume.pdf")
SITE_PDF = os.path.join(SITE, "assets", "files", "Margolin_Resume.pdf")
SITE_MD = os.path.join(SITE, "resume", "index.md")

# ----------------------------------------------------------------------------
# CONTENT  (**bold** markup is honoured in every renderer)
# ----------------------------------------------------------------------------
CONTENT = {
    "name": 'Elizabeth "Eli" Margolin',
    "contact": [
        ("ecmargo@seas.upenn.edu", "mailto:ecmargo@seas.upenn.edu"),
        ("ecmargo.github.io", "https://ecmargo.github.io"),
        ("github.com/ecmargo", "https://github.com/ecmargo"),
        ("Google Scholar", "https://scholar.google.com/citations?user=HzZ82ZMAAAAJ&hl=en"),
    ],
    "summary": (
        "PhD candidate at the University of Pennsylvania building zero-knowledge proof systems at the "
        "intersection of programming languages and cryptography, with publications at IEEE S&P, USENIX "
        "Security, and SOSP. Three years of security engineering at Meta. Current work applies these tools "
        "to AI accountability, including zkTLS provenance for LLM responses."
    ),
    "skills": [
        ("Cryptography",
         "zero-knowledge proofs; secure two-party and multi-party computation; homomorphic encryption; "
         "differential privacy."),
        ("Programming",
         "Rust, C++, Python; interactive theorem proving (Rocq, Lean)."),
        ("Security",
         "threat modeling; abuse and insider-threat detection; security-sensitive code review."),
    ],
    "experience": [
        {
            "org": "University of Pennsylvania",
            "role": "PhD Researcher, Cryptography & Programming Languages",
            "dates": "2021–present",
            "bullets": [
                "Designed and built zero-knowledge proof systems for formal languages and zkTLS protocols that "
                "let users prove facts about private data and real web sessions; owned the work from problem "
                "formulation and protocol design through open-source implementation and evaluation.",
                "Designed and built a planner that automatically compiles differentially private analytics "
                "queries into secure multi-party and homomorphic computations at billion-device scale.",
                "Work published at IEEE S&P, USENIX Security, and SOSP; two further papers in submission.",  # SITE_VARIANT
                "Advised by Sebastian Angel. Dissertation: *Zero Knowledge Proofs of Formal Languages and "
                "Their Applications*.",
            ],
        },
        {
            "org": "Brave Software",
            "role": "Research Intern",
            "dates": "Summer 2024",
            "bullets": [
                "Designed zero-knowledge protocols for privacy-preserving fraud detection over authenticated "
                "web traffic, taking research designs through to deployable implementations.",
                "Built and hardened cryptographic tooling to interoperate with modern TLS stacks.",
            ],
        },
        {
            "org": "Meta Platforms",
            "role": "Security Engineer",
            "dates": "2019–2021 (Contingent Worker, 2022)",
            "bullets": [
                "Performed threat modeling and security analysis to design and own company-wide insider-threat "
                "and abuse-detection tooling for security and legal investigations teams.",
                "Scaled alert-processing pipelines through automation and led security and privacy reviews "
                "with Legal and Policy, raising throughput while cutting false positives and data footprint.",
            ],
        },
    ],
    "publications": [
        {"title": "Coral: Fast Succinct Non-Interactive Zero-Knowledge CFG Proofs",
         "authors": "S. Angel, S. Celi, E. Margolin*, P. Mishra, M. Sander, J. Woods",
         "venue": "IEEE S&P 2026",
         "links": [("paper", "https://eprint.iacr.org/2025/1420"), ("code", "https://github.com/eniac/coral")]},
        {"title": "Reef: Fast Succinct Non-Interactive Zero-Knowledge Regex Proofs",
         "authors": "S. Angel, E. Ioannidis, E. Margolin*, S. Setty, J. Woods",
         "venue": "USENIX Security 2024",
         "links": [("paper", "https://eprint.iacr.org/2023/1886"), ("code", "https://github.com/eniac/Reef")]},
        {"title": "Arboretum: A Planner for Large-Scale Federated Analytics with Differential Privacy",
         "authors": "E. Margolin, K. Newatia, E. Roth, T. Luo, A. Haeberlen",
         "venue": "SOSP 2023",
         "links": [("paper", "/assets/pdf/arboretum-sosp2023.pdf"), ("code", "https://github.com/ecmargo/arboretum")]},
        {"title": "Weasel: Zero Knowledge Proofs of WASM Compilation",
         "authors": "J. Woods, E. Margolin, E. Ioannidis, S. Angel, P. Mishra",
         "venue": "In submission",
         "links": [], "site": False},
        {"title": "Surf: Bringing zkTLS to the Modern Web",
         "authors": "S. Angel, S. Celi, E. Margolin*",
         "venue": "In submission",
         "links": [], "site": False},
    ],
    "pub_note": "* Authors listed in alphabetical order.",
    "education": [
        ("University of Pennsylvania", "PhD, Computer and Information Science", "In progress"),
        ("Duke University", "MS, Economics and Computation", ""),
        ("Stanford University", "BA, Political Science, with Honors in International Security", ""),
    ],
    "activities": [
        ("Team USA", "National Team Athlete (Rowing); USRowing Athlete Council", "2025–present",
         ["1st place, 2026 World Rowing Cup III, Lucerne.",
          "A Finalist, 2025 World Rowing Championships, Shanghai."]),
    ],
}

SITE_VARIANTS = {
    "Work published at IEEE S&P, USENIX Security, and SOSP; two further papers in submission.":
    "Work published at IEEE S&P, USENIX Security, and SOSP.",
}


def site_content():
    """CONTENT for the website: omits publications marked site: False and swaps site-variant strings."""
    import copy
    c = copy.deepcopy(CONTENT)
    c["publications"] = [pub for pub in c["publications"] if pub.get("site", True)]
    for e in c["experience"]:
        e["bullets"] = [SITE_VARIANTS.get(b, b) for b in e["bullets"]]
    return c


BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
MARK_RE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")


def segments(text):
    """Split '**bold**' / '*italic*' markup into [(str, is_bold, is_italic)]."""
    out, pos = [], 0
    for m in MARK_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False))
        if m.group(1) is not None:
            out.append((m.group(1), True, False))
        else:
            out.append((m.group(2), False, True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False))
    return out


# ----------------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------------
def build_docx(path, C=CONTENT):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    BLUE = RGBColor(0x1F, 0x4E, 0x79)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    width_in = 8.5 - 1.4

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    lb = doc.styles["List Bullet"]
    lb.paragraph_format.space_after = Pt(1)

    def add_hyperlink(paragraph, text, url, bold=False):
        part = paragraph.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "1F4E79"); rpr.append(color)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        if bold:
            rpr.append(OxmlElement("w:b"))
        run.append(rpr)
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
        run.append(t)
        link.append(run)
        paragraph._p.append(link)

    def add_runs(paragraph, text, size=None, italic=False):
        for s, b, i in segments(text):
            r = paragraph.add_run(s)
            r.bold = b
            r.italic = italic or i
            if size:
                r.font.size = Pt(size)

    def header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title.upper())
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = BLUE
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F4E79")):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        ppr.append(pbdr)

    def entry_line(left, right, left_bold=True, right_italic=True):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(4)
        for s, b, i in segments(left):
            r = p.add_run(s); r.bold = b or left_bold; r.italic = i
        if right:
            r = p.add_run("\t" + right); r.italic = right_italic
        return p

    def bullet(text, size=None):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        p.paragraph_format.space_after = Pt(0)
        add_runs(p, text, size=size)
        return p

    # Header
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(C["name"]); r.bold = True; r.font.size = Pt(20)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (txt, url) in enumerate(C["contact"]):
        if i:
            p.add_run("  ·  ")
        add_hyperlink(p, txt, url)

    header("Summary")
    p = doc.add_paragraph(); add_runs(p, C["summary"])

    header("Technical Skills")
    for label, body in C["skills"]:
        p = doc.add_paragraph()
        r = p.add_run(label + ": "); r.bold = True
        add_runs(p, body)

    header("Experience")
    for e in C["experience"]:
        entry_line(f'{e["org"]} — {e["role"]}', e["dates"])
        for b in e["bullets"]:
            bullet(b)

    header("Selected Publications")
    for pub in C["publications"]:
        entry_line(pub["title"], pub["venue"])
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(pub["authors"]); r.font.size = Pt(9.5)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2)
    r = p.add_run(C["pub_note"]); r.italic = True; r.font.size = Pt(9)

    header("Education")
    for org, deg, when in C["education"]:
        entry_line(f"{org} — {deg}", when)

    header("Selected Activities")
    for org, role, when, desc in C["activities"]:
        entry_line(f"{org} — {role}", when)
        for d in desc:
            bullet(d)

    doc.save(path)
    return path
    return path


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
def build_pdf(path, C=CONTENT):
    from fpdf import FPDF

    FONT_DIR = "/System/Library/Fonts/Supplemental"
    pdf = FPDF(format="letter", unit="in")
    pdf.set_margins(0.7, 0.5, 0.7)
    pdf.set_auto_page_break(auto=True, margin=0.45)
    pdf.add_font("Arial", "", os.path.join(FONT_DIR, "Arial.ttf"))
    pdf.add_font("Arial", "B", os.path.join(FONT_DIR, "Arial Bold.ttf"))
    pdf.add_font("Arial", "I", os.path.join(FONT_DIR, "Arial Italic.ttf"))
    pdf.add_page()
    BLUE = (0x1F, 0x4E, 0x79)
    BASE = 10
    LH = 0.163  # line height in inches at 10pt
    W = pdf.epw

    def md(text):
        # fpdf markdown: **bold**, __italic__
        return re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"__\1__", text)

    def header(title):
        pdf.ln(0.075)
        pdf.set_font("Arial", "B", 10.5); pdf.set_text_color(*BLUE)
        pdf.cell(W, LH, title.upper(), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.set_draw_color(*BLUE); pdf.set_line_width(0.007)
        pdf.line(pdf.l_margin, y, pdf.l_margin + W, y)
        pdf.ln(0.05)
        pdf.set_text_color(0, 0, 0)

    def para(text, size=BASE, indent=0.0, italic=False, lh=None):
        pdf.set_font("Arial", "I" if italic else "", size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(W - indent, lh or LH, md(text), markdown=True, new_x="LMARGIN", new_y="NEXT")

    def entry_line(left, right):
        pdf.ln(0.05)
        pdf.set_font("Arial", "B", BASE)
        rw = 0
        if right:
            pdf.set_font("Arial", "I", BASE); rw = pdf.get_string_width(right) + 0.05
        pdf.set_font("Arial", "B", BASE)
        pdf.cell(W - rw, LH, left, new_x="RIGHT", new_y="TOP")
        if right:
            pdf.set_font("Arial", "I", BASE)
            pdf.cell(rw, LH, right, align="R", new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.ln(LH)

    def bullet(text):
        pdf.set_font("Arial", "", BASE)
        x0 = pdf.l_margin + 0.08
        pdf.set_x(x0); pdf.cell(0.12, LH, "•", new_x="RIGHT", new_y="TOP")
        pdf.multi_cell(W - 0.20, LH, md(text), markdown=True, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.015)

    # Header
    pdf.set_font("Arial", "B", 19)
    pdf.cell(W, 0.32, C["name"], align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", BASE); pdf.set_text_color(*BLUE)
    parts = C["contact"]
    sep = "   ·   "
    total = sum(pdf.get_string_width(t) for t, _ in parts) + pdf.get_string_width(sep) * (len(parts) - 1)
    pdf.set_x(pdf.l_margin + (W - total) / 2)
    for i, (txt, url) in enumerate(parts):
        if i:
            pdf.set_text_color(0, 0, 0); pdf.cell(pdf.get_string_width(sep), LH, sep, new_x="RIGHT", new_y="TOP")
            pdf.set_text_color(*BLUE)
        pdf.cell(pdf.get_string_width(txt), LH, txt, link=url, new_x="RIGHT", new_y="TOP")
    pdf.ln(LH); pdf.set_text_color(0, 0, 0)

    header("Summary")
    para(C["summary"])

    header("Technical Skills")
    for label, body in C["skills"]:
        para(f"**{label}:** {body}")

    header("Experience")
    for e in C["experience"]:
        entry_line(f'{e["org"]} — {e["role"]}', e["dates"])
        for b in e["bullets"]:
            bullet(b)

    header("Selected Publications")
    for pub in C["publications"]:
        entry_line(pub["title"], pub["venue"])
        para(pub["authors"], size=9.5, indent=0.15, lh=0.165)
    pdf.ln(0.02); para(C["pub_note"], size=9, italic=True, lh=0.16)

    header("Education")
    for org, deg, when in C["education"]:
        entry_line(f"{org} — {deg}", when)

    header("Selected Activities")
    for org, role, when, desc in C["activities"]:
        entry_line(f"{org} — {role}", when)
        for d in desc:
            bullet(d)

    pages = pdf.page_no()
    print(f"pdf: ended on page {pages} at y={pdf.get_y():.2f} in (page height 11.0, bottom margin 0.45)")
    pdf.output(path)
    return pages


# ----------------------------------------------------------------------------
# Jekyll resume page
# ----------------------------------------------------------------------------
def build_site_md(path, C=CONTENT):
    def esc(s):
        return s.replace("&", "&amp;").replace('"', '\\"')

    def html_bold(s):
        s = BOLD_RE.sub(r"<strong>\1</strong>", s)
        return re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", s)

    lines = ["---", "layout: list", "title: Eli Margolin - Resume",
             "download: /assets/files/Margolin_Resume.pdf", "sections:"]

    lines += ["  - label: Summary", "    items:", f'      - body: "{esc(C["summary"])}"']

    lines += ["  - label: Technical Skills", "    items:"]
    for label, body in C["skills"]:
        lines += [f'      - title: "{esc(label)}"', f'        body: "{esc(body)}"']

    lines += ["  - label: Experience", "    items:"]
    for e in C["experience"]:
        ul = "<ul>" + "".join(f"<li>{html_bold(esc(b))}</li>" for b in e["bullets"]) + "</ul>"
        lines += [f'      - title: "{esc(e["org"])} — {esc(e["role"])}"',
                  f'        meta: "{esc(e["dates"])}"',
                  f'        body: "{ul}"']

    lines += ["  - label: Selected Publications", "    items:"]
    for i, pub in enumerate(C["publications"]):
        body = esc(pub["authors"]) + "."
        if i == len(C["publications"]) - 1:
            body += f' <em>{esc(C["pub_note"])}</em>'
        lines += [f'      - title: "{esc(pub["title"])}"',
                  f'        meta: "{esc(pub["venue"])}"',
                  f'        body: "{body}"']
        if pub["links"]:
            lines.append("        links:")
            for text, url in pub["links"]:
                lines += [f"          - text: {text}", f"            url: {url}"]

    lines += ["  - label: Education", "    items:"]
    for org, deg, when in C["education"]:
        lines.append(f'      - title: "{esc(org)} — {esc(deg)}"')
        if when:
            lines.append(f'        meta: "{esc(when)}"')

    lines += ["  - label: Selected Activities", "    items:"]
    for org, role, when, desc in C["activities"]:
        lines += [f'      - title: "{esc(org)} — {esc(role)}"',
                  f'        meta: "{esc(when)}"']
        if desc:
            lines.append('        body: "<ul>' + "".join(f"<li>{esc(d)}</li>" for d in desc) + '</ul>"')

    lines += ["---", ""]
    with open(path, "w") as f:
        f.write("\n".join(lines))
    return path


if __name__ == "__main__":
    build_docx(DOCX_OUT)
    pages = build_pdf(PDF_OUT)
    SC = site_content()
    site_pages = build_pdf(SITE_PDF, SC)
    build_site_md(SITE_MD, SC)
    print(f"docx -> {DOCX_OUT}")
    print(f"pdf  -> {PDF_OUT} ({pages} page{'s' if pages != 1 else ''})")
    print(f"site -> {SITE_PDF} ({site_pages} page{'s' if site_pages != 1 else ''}), {SITE_MD}")
    if pages != 1 or site_pages != 1:
        print("WARNING: a PDF is not one page", file=sys.stderr)
        sys.exit(1)
